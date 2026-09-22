"""
Reading a job-application tracker someone kept in a spreadsheet: which sheet,
which header row, which column is which field, and each row as application
fields plus what is wrong with it. Pure — no database — so every rule here is
testable on a file alone. Importing the result is application_import.py.

Built against a real tracker, so it expects real-tracker mess: "Location"
holding Remote/Onsite, a link column holding company names, outcomes kept in a
"Next Action" column, a "Read Me" sheet beside the data.
"""
import csv
import io
import re
import string
from datetime import date, datetime, timedelta

from src.jobpool.posting import employment_text

MAX_BYTES = 5_000_000
MAX_ROWS = 2_000

FIELDS = {"title": "Job title", "company": "Company", "url": "Job link", "location": "Location",
          "workplace": "Remote / onsite", "applied_at": "Date applied", "status": "Status",
          "outcome": "Outcome (moves the status on)", "employment_type": "Job type", "salary": "Salary",
          "source": "Where you found it", "notes": "Notes"}
# Header text -> field. Several columns may feed notes; every other field takes one.
HEADERS = {
    "title": "job title|title|role|position|job|job role|position title|job name|vacancy|job position|role title|opening",
    "company": "company|company name|employer|organization|organisation|org|firm|hiring company|business",
    "url": "url|link|job url|job link|posting url|job posting url|job posting|posting|application link|apply link|"
           "job ad|advert|listing|listing url",
    "location": "location|city|place|country|job location|office|based in",
    "workplace": "remote|work mode|workplace|work type|arrangement|remote onsite|remote hybrid onsite|work setting|"
                 "work model|on site remote|remote or onsite",
    "applied_at": "date applied|applied|application date|applied on|date|applied date|submitted|date submitted|"
                  "submission date|applied at|date of application|when applied",
    "status": "status|stage|application status|state|current status|pipeline stage",
    "outcome": "outcome|result|response|decision|reply|feedback",
    "employment_type": "job type|type|employment type|contract|contract type|employment|commitment",
    "salary": "salary|pay|compensation|salary range|rate|expected salary|package|stipend",
    "source": "source|found on|where found|platform|board|job board|site|found via|channel|how found",
    "notes": "notes|note|comments|comment|remarks|details|next action|next step|next steps|follow up|contact|"
             "recruiter|resume|cv|cover letter|resume cover letter|resume cover letter folder|priority",
}
_HEADER_FIELD = {h: field for field, names in HEADERS.items() for h in names.split("|")}

STATUS_WORDS = {
    "saved": "not applied|saved|wishlist|wish list|bookmarked|interested|to apply|planning|planned|draft|considering",
    "applied": "applied|submitted|sent|application sent|applied online",
    "screening": "screening|phone screen|recruiter call|hr call|screen|assessment|test|take home|coding test|"
                 "online assessment|in review|under review|reviewing|shortlisted",
    "interview": "interview|interviewing|interviews|first interview|second interview|final interview|final round|"
                 "onsite|technical interview|panel",
    "offer": "offer|offered|offer received|accepted|success|successful|hired|got the job",
    "rejected": "rejected|declined|not selected|unsuccessful|turned down|rejection|regret|closed",
    "withdrawn": "withdrawn|withdrew|declined offer|not interested|cancelled|canceled|dropped",
}
_STATUS = {w: status for status, words in STATUS_WORDS.items() for w in words.split("|")}
# Words that describe waiting, not a move: kept as a note, status unchanged.
_WAITING = {"no response", "await response", "awaiting response", "waiting", "pending", "ghosted", "no reply", "none"}
PIPELINE = ["saved", "applied", "screening", "interview", "offer"]
WORKPLACES = {"remote": "remote", "wfh": "remote", "work from home": "remote", "yes": "remote", "true": "remote",
              "hybrid": "hybrid", "onsite": "onsite", "on site": "onsite", "in office": "onsite", "office": "onsite",
              "no": "onsite", "false": "onsite"}
SOURCES = {"linkedin": "linkedin", "indeed": "indeed", "glassdoor": "glassdoor", "company site": "company_site",
           "company website": "company_site", "careers page": "company_site", "website": "company_site",
           "direct": "company_site", "referral": "referral", "recruiter": "recruiter", "job fair": "job_fair",
           "career fair": "job_fair", "brightermonday": "brightermonday", "fuzu": "fuzu", "myjobmag": "myjobmag",
           "other": "other"}


class InvalidSpreadsheet(ValueError):
    pass


def _words(value) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).split())


def cell_text(value) -> str:
    """A cell as one line of text: 3.0 reads "3", None reads ""."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


# ------------------------------------------------------------------ reading

def read_table(data: bytes, file_name: str) -> tuple[str | None, list[str], list[list]]:
    """(sheet used, every sheet's name, its rows). xlsx and csv; the old binary .xls is not read."""
    if len(data) > MAX_BYTES:
        raise InvalidSpreadsheet("That file is over 5 MB. Upload just the sheet with your applications.")
    if not data.strip():
        raise InvalidSpreadsheet("That file is empty.")
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise InvalidSpreadsheet("That is the old Excel .xls format. Save it as .xlsx or .csv and upload again.")
    if data[:2] == b"PK":
        return _read_xlsx(data)
    if file_name.lower().endswith((".xlsx", ".xlsm")):
        raise InvalidSpreadsheet("That file couldn't be read as a spreadsheet. Upload an .xlsx or .csv file.")
    return None, [], _read_csv(data)


def _read_xlsx(data: bytes):
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as err:
        raise InvalidSpreadsheet("That file couldn't be read as a spreadsheet. Upload an .xlsx or .csv file.") from err
    sheets = {ws.title: [list(r) for r, _ in zip(ws.iter_rows(values_only=True), range(MAX_ROWS + 20))]
              for ws in book.worksheets}
    book.close()
    # The sheet whose header reads most like an applications tracker.
    best = max(sheets, key=lambda name: find_header(sheets[name])[1])
    return best, list(sheets), sheets[best]


def _read_csv(data: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row, _ in zip(csv.reader(io.StringIO(text), dialect), range(MAX_ROWS + 20))]


def find_header(rows: list[list]) -> tuple[int | None, int]:
    """(index, recognised cells) of the row among the first ten that reads most like a header."""
    best, hits = None, 0
    for i, row in enumerate(rows[:10]):
        found = sum(1 for cell in row if _words(cell) in _HEADER_FIELD)
        if found > hits:
            best, hits = i, found
    return (best, hits) if hits >= 2 else (None, 0)


# ------------------------------------------------------------------ mapping

def _letter(i: int) -> str:
    return string.ascii_uppercase[i] if i < 26 else string.ascii_uppercase[i // 26 - 1] + string.ascii_uppercase[i % 26]


def _outcome_like(values) -> bool:
    words = [_words(v) for v in values if cell_text(v)]
    return bool(words) and sum(1 for w in words if w in _STATUS or w in _WAITING) >= 0.6 * len(words)


def map_columns(header: list, rows: list[list], chosen: dict | None = None) -> list[dict]:
    """
    Each column's field and how it was decided: "header", "values" (read from
    the cells), or "you" (the user's correction). Unmapped columns stay listed
    with their samples, so the user can see what is being ignored.
    """
    width = max([len(header), *(len(r) for r in rows[:50])], default=0)
    columns, taken = [], set()
    for i in range(width):
        name = cell_text(header[i]) if i < len(header) else ""
        values = [r[i] for r in rows[:200] if i < len(r) and cell_text(r[i])]
        field, how = _HEADER_FIELD.get(_words(name)), "header"
        # A "Next Action" or "Response" column full of outcomes is an outcome column.
        if field in ("notes", "outcome", None) and values and _outcome_like(values):
            field, how = ("outcome" if field or "status" in taken else "status"), "values"
        if field is None and values and sum(1 for v in values if _looks_like_link(cell_text(v))) >= 0.6 * len(values):
            field, how = "url", "values"
        if field and field != "notes" and field in taken:
            field, how = None, None  # a second column cannot also be the job title
        columns.append({"index": i, "letter": _letter(i), "header": name or f"Column {_letter(i)}",
                        "field": field, "how": how if field else None,
                        "samples": [cell_text(v)[:60] for v in values[:3]]})
        if field:
            taken.add(field)
    for i, field in (chosen or {}).items():
        i = int(i)
        if 0 <= i < len(columns):
            if field and field != "notes":
                for other in columns:
                    if other["field"] == field:
                        other.update(field=None, how=None)
            columns[i].update(field=field or None, how="you" if field else None)
    return columns


# ------------------------------------------------------------------- values

_DMY = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})$")


def date_order(values) -> tuple[str, bool]:
    """("dmy" | "mdy", ambiguous): read from the whole column, since "03/04" alone cannot say."""
    firsts, seconds = [], []
    for v in values:
        m = _DMY.match(cell_text(v))
        if m:
            firsts.append(int(m.group(1)))
            seconds.append(int(m.group(2)))
    if any(f > 12 for f in firsts):
        return "dmy", False
    if any(s > 12 for s in seconds):
        return "mdy", False
    return "dmy", bool(firsts)


def parse_date(value, order: str = "dmy") -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:
        return date(1899, 12, 30) + timedelta(days=int(value))   # an Excel serial number
    text = cell_text(value)
    if not text:
        return None
    m = _DMY.match(text)
    if m:
        a, b, year = int(m.group(1)), int(m.group(2)), int(m.group(3))
        day, month = (a, b) if order == "dmy" else (b, a)
        year = year + 2000 if year < 100 else year
    else:
        try:
            return datetime.fromisoformat(text.replace("/", "-")[:19]).date()
        except ValueError:
            pass
        cleaned = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", text.replace(",", " "))
        for pattern in ("%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y", "%b %Y", "%B %Y"):
            try:
                return datetime.strptime(" ".join(cleaned.split()), pattern).date()
            except ValueError:
                continue
        return None
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _looks_like_link(text: str) -> bool:
    return bool(re.match(r"^(https?://)?(www\.)?[\w-]+(\.[\w-]+)+(/\S*)?$", text.strip(), re.I))


def link(text: str) -> str | None:
    text = text.strip()
    if not _looks_like_link(text):
        return None
    return text if re.match(r"^https?://", text, re.I) else f"https://{text}"


def _status(word: str) -> str | None:
    return _STATUS.get(_words(word))


def _furthest(status: str, outcome: str | None) -> str:
    if outcome in ("rejected", "withdrawn"):
        return outcome
    if outcome in PIPELINE and status in PIPELINE and PIPELINE.index(outcome) > PIPELINE.index(status):
        return outcome
    return status


def normalise_row(values: list, columns: list[dict], order: str, today: date) -> tuple[dict, list[dict]]:
    """One sheet row -> application fields, plus what is wrong with it."""
    cell = {c["field"]: values[c["index"]] for c in columns
            if c["field"] and c["field"] != "notes" and c["index"] < len(values)}
    notes = [f"{c['header']}: {cell_text(values[c['index']])}" for c in columns
             if c["field"] == "notes" and c["index"] < len(values) and cell_text(values[c["index"]])]
    issues = []

    def flag(field, message, level="warning"):
        issues.append({"field": field, "message": message, "level": level})

    row = {"title": cell_text(cell.get("title"))[:300] or None, "company": cell_text(cell.get("company"))[:300] or None}

    raw_link = cell_text(cell.get("url"))
    row["url"] = link(raw_link)[:2048] if link(raw_link) else None
    if raw_link and not row["url"]:
        notes.insert(0, f"Link: {raw_link}")
        flag("url", "Not a web link — kept in notes")

    place = cell_text(cell.get("location"))
    arrangement = WORKPLACES.get(_words(cell.get("workplace"))) if cell.get("workplace") is not None else None
    if _words(place) in WORKPLACES and _words(place) not in ("yes", "no", "true", "false"):
        arrangement, place = arrangement or WORKPLACES[_words(place)], ""  # "Remote" is an arrangement, not a place
    elif not arrangement:
        arrangement = next((w for word, w in WORKPLACES.items()
                            if len(word) > 3 and re.search(rf"\b{word}\b", _words(place))), None)
    row["location"], row["workplace"] = place[:300] or None, arrangement

    stated = cell_text(cell.get("status"))
    status = _status(stated) if stated else None
    if stated and status is None and _words(stated) not in _WAITING:
        flag("status", f"Status \"{stated}\" not recognised — imported as Applied")
    if stated and _words(stated) in _WAITING:
        notes.append(f"Status: {stated}")
    outcome_text = cell_text(cell.get("outcome"))
    outcome = _status(outcome_text) if outcome_text else None
    if outcome_text and outcome is None:
        notes.append(f"Outcome: {outcome_text}")

    applied = parse_date(cell.get("applied_at"), order)
    if cell.get("applied_at") not in (None, "") and applied is None:
        flag("applied_at", f"Couldn't read the date \"{cell_text(cell.get('applied_at'))}\"")
    elif applied and (applied > today + timedelta(days=1) or applied.year < 2000):
        flag("applied_at", f"Date {applied.isoformat()} is not a plausible application date — left empty")
        applied = None
    row["status"] = _furthest(status or "applied", outcome)
    if row["status"] == "saved":
        applied = None  # a saved job has not been applied to, whatever the sheet says
    row["applied_at"] = applied.isoformat() if applied else None

    row["employment_type"] = (employment_text(cell_text(cell.get("employment_type"))) or None)
    row["salary"] = cell_text(cell.get("salary"))[:200] or None
    source = _words(cell.get("source"))
    row["source"] = SOURCES.get(source, source.replace(" ", "_")[:50] or None) if source else None
    row["notes"] = "\n".join(notes)[:2000] or None

    if not row["title"]:
        flag("title", "No job title — can't import this row", "error")
    if not row["company"]:
        flag("company", "No company")
    if row["status"] != "saved" and not row["applied_at"] and not any(i["field"] == "applied_at" for i in issues):
        flag("applied_at", "No application date — it will show as unknown")
    return row, issues
